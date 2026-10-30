import docx
from docx.shared import Inches, Pt, RGBColor

doc = docx.Document()

for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

with open('panduan-lengkap-hermes.md', 'r', encoding='utf-8') as f:
    lines = f.readlines()

for line in lines:
    line_s = line.strip()
    if line_s.startswith('# '):
        p = doc.add_paragraph()
        run = p.add_run(line_s[2:])
        run.bold = True
        run.font.size = Pt(20)
        run.font.color.rgb = RGBColor(31, 78, 121)
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
    elif line_s.startswith('## '):
        p = doc.add_paragraph()
        run = p.add_run(line_s[3:])
        run.bold = True
        run.font.size = Pt(15)
        run.font.color.rgb = RGBColor(46, 117, 182)
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
    elif line_s.startswith('### '):
        p = doc.add_paragraph()
        run = p.add_run(line_s[4:])
        run.bold = True
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(89, 89, 89)
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
    elif line_s.startswith('- ') or line_s.startswith('* '):
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(line_s[2:])
        p.paragraph_format.space_after = Pt(3)
    elif line_s.startswith('```'):
        continue
    elif line_s == '---':
        doc.add_paragraph('___________________________________________________').paragraph_format.space_after = Pt(6)
    elif line_s != '':
        p = doc.add_paragraph()
        p.add_run(line_s)
        p.paragraph_format.space_after = Pt(4)

output_path = 'Panduan_Setup_Dan_Penggunaan_Hermes_SymbioNusa.docx'
doc.save(output_path)
print('SUCCESS:', output_path)
